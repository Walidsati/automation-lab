#!/usr/bin/env python3
"""Inspect the xlsxwriter-generated Excel file for fills."""

from pathlib import Path
from openpyxl import load_workbook

FILE = Path(__file__).resolve().parent / "data" / "sales_report_xw.xlsx"
wb = load_workbook(FILE)

for sheet_name in ("By Region", "By Month"):
    ws = wb[sheet_name]
    print(f"\n=== {sheet_name} ===")
    for row in ws.iter_rows():
        for cell in row:
            fill = None
            if cell.fill and cell.fill.patternType:
                fill = cell.fill.fgColor.rgb
            marker = ""
            if fill and fill not in ("00000000", "FFFFFFFF", None):
                marker = f"  ← fill={fill}"
            print(f"{cell.coordinate}: {cell.value!r}{marker}")
