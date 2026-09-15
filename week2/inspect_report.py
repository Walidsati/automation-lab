#!/usr/bin/env python3
"""Inspect the generated Excel file."""

from pathlib import Path
from openpyxl import load_workbook

SCRIPT_DIR = Path(__file__).resolve().parent
FILE = SCRIPT_DIR / "data" / "sales_report.xlsx"

wb = load_workbook(FILE)

for ws in wb.worksheets:
    print(f"\n=== Sheet: {ws.title} ===")
    print(f"Dimensions: {ws.dimensions}, max_row={ws.max_row}, max_col={ws.max_column}")
    for row_idx in range(1, ws.max_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            fill = cell.fill
            fg = getattr(fill.fgColor, "rgb", None) if fill and fill.patternType else None
            marker = ""
            if fg and fg not in ("00000000", "FFFFFFFF", None):
                marker = f"  ← fill={fg}"
            print(f"  {cell.coordinate}: {cell.value!r}{marker}")
