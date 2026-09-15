#!/usr/bin/env python3
"""Generate a formatted Excel report from sales.csv."""

from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ---- Paths ----
SCRIPT_DIR = Path(__file__).resolve().parent
INPUT = SCRIPT_DIR / "data" / "sales.csv"
OUTPUT = SCRIPT_DIR / "data" / "sales_report.xlsx"

# ---- Colors & Styles (reused everywhere) ----
HEADER_BG = "4472C4"       # blue
HEADER_FG = "FFFFFF"       # white
TOTAL_BG = "FFF2CC"        # light yellow
BEST_BG = "C6EFCE"         # light green

header_font = Font(bold=True, color=HEADER_FG, size=12)
header_fill = PatternFill("solid", fgColor=HEADER_BG)
header_align = Alignment(horizontal="center", vertical="center")

total_font = Font(bold=True)
total_fill = PatternFill("solid", fgColor=TOTAL_BG)

best_fill = PatternFill("solid", fgColor=BEST_BG)

thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

money_format = '"$"#,##0.00'


def style_header(ws, row_idx, num_cols):
    """Apply header styling to a row."""
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border


def autosize_columns(ws, min_width=8, max_width=30):
    """Set column widths based on longest value in each column."""
    for col_idx, col in enumerate(ws.iter_cols(), start=1):
        max_len = 0
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        width = min(max(max_len + 2, min_width), max_width)
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def write_dataframe(ws, df, start_row=1, money_columns=None):
    """Write a DataFrame starting at the given row, with styled header."""
    money_columns = money_columns or []

    # Header
    for col_idx, col_name in enumerate(df.columns, start=1):
        ws.cell(row=start_row, column=col_idx, value=col_name)
    style_header(ws, start_row, len(df.columns))

    # Data rows
    for row_idx, (_, row) in enumerate(df.iterrows(), start=start_row + 1):
        for col_idx, col_name in enumerate(df.columns, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row[col_name])
            cell.border = thin_border
            if col_name in money_columns:
                cell.number_format = money_format


def main():
    # ---- 1. Load & prepare ----
    df = pd.read_csv(INPUT)
    df["date"] = pd.to_datetime(df["date"])
    df["revenue"] = (df["units"] * df["unit_price"]).round(2)

    # ---- 2. Prep summaries ----
    by_region = (
        df.groupby("region")
        .agg(units=("units", "sum"), revenue=("revenue", "sum"))
        .round(2)
        .sort_values("revenue", ascending=False)
        .reset_index()
    )
    total_row = pd.DataFrame([{
        "region": "TOTAL",
        "units": by_region["units"].sum(),
        "revenue": by_region["revenue"].sum(),
    }])
    by_region_with_total = pd.concat([by_region, total_row], ignore_index=True)

    df["month"] = df["date"].dt.strftime("%Y-%m")
    by_month = (
        df.groupby("month")
        .agg(units=("units", "sum"), revenue=("revenue", "sum"))
        .round(2)
        .reset_index()
    )

    # ---- 3. Build workbook ----
    wb = Workbook()

    # Sheet 1: Raw Data
    ws1 = wb.active
    ws1.title = "Raw Data"
    write_dataframe(ws1, df[["date", "region", "product", "units", "unit_price", "revenue"]],
                    money_columns=["unit_price", "revenue"])
    ws1.freeze_panes = "A2"     # freeze header row when scrolling
    autosize_columns(ws1)

    # Sheet 2: By Region
    ws2 = wb.create_sheet("By Region")
    write_dataframe(ws2, by_region_with_total, money_columns=["revenue"])
    # Style the total row
    last_row = ws2.max_row
    for col in range(1, 4):
        cell = ws2.cell(row=last_row, column=col)
        cell.font = total_font
        cell.fill = total_fill
    autosize_columns(ws2)

    # Sheet 3: By Month
    ws3 = wb.create_sheet("By Month")
    write_dataframe(ws3, by_month, money_columns=["revenue"])
    # Highlight the best month
    best_revenue = by_month["revenue"].max()
    for row_idx in range(2, ws3.max_row + 1):
        if ws3.cell(row=row_idx, column=3).value == best_revenue:
            for col in range(1, 4):
                ws3.cell(row=row_idx, column=col).fill = best_fill
    autosize_columns(ws3)

    # ---- 4. Save ----
    wb.save(OUTPUT)
    print(f"Saved: {OUTPUT}")
    print(f"Sheets: {[s.title for s in wb.worksheets]}")


if __name__ == "__main__":
    main()
